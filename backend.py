"""
GST Inward HSN/SAC Fetcher — backend logic
Handles: SQLite storage, e-invoice JSON decoding, Excel report generation.
No network calls in V1 — everything operates on locally supplied files.
"""
import json
import base64
import sqlite3
import os
import re
import datetime
from pathlib import Path

import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

APP_DIR = Path(os.environ.get("APPDATA", str(Path.home()))) / "GSTInwardHSNFetcher"
APP_DIR.mkdir(parents=True, exist_ok=True)
DB_PATH = APP_DIR / "data.db"

HEADER_FILL = "1A5FA8"


def _b64pad(s):
    return s + "=" * (-len(s) % 4)


def get_conn():
    conn = sqlite3.connect(str(DB_PATH))
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = get_conn()
    cur = conn.cursor()
    cur.executescript(
        """
        CREATE TABLE IF NOT EXISTS settings (
            key TEXT PRIMARY KEY,
            value TEXT
        );
        CREATE TABLE IF NOT EXISTS gstins (
            gstin TEXT PRIMARY KEY,
            label TEXT
        );
        CREATE TABLE IF NOT EXISTS invoices (
            irn TEXT PRIMARY KEY,
            gstin TEXT,
            ack_no TEXT,
            ack_dt TEXT,
            status TEXT,
            doc_type TEXT,
            invoice_no TEXT,
            invoice_date TEXT,
            supplier_gstin TEXT,
            supplier_name TEXT,
            buyer_gstin TEXT,
            buyer_name TEXT,
            taxable_value REAL,
            igst REAL,
            cgst REAL,
            sgst REAL,
            cess REAL,
            total_value REAL,
            period TEXT
        );
        CREATE TABLE IF NOT EXISTS items (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            irn TEXT,
            sl_no TEXT,
            hsn TEXT,
            description TEXT,
            is_service TEXT,
            qty REAL,
            unit TEXT,
            unit_price REAL,
            taxable_amount REAL,
            gst_rate REAL,
            igst_amt REAL,
            cgst_amt REAL,
            sgst_amt REAL,
            cess_amt REAL,
            total_item_value REAL
        );
        CREATE TABLE IF NOT EXISTS exports (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            filename TEXT,
            filepath TEXT,
            period_from TEXT,
            period_to TEXT,
            gstin TEXT,
            invoice_count INTEGER,
            generated_at TEXT
        );
        """
    )
    conn.commit()
    conn.close()


# ---------- settings ----------

def get_setting(key, default=None):
    conn = get_conn()
    row = conn.execute("SELECT value FROM settings WHERE key=?", (key,)).fetchone()
    conn.close()
    return row["value"] if row else default


def set_setting(key, value):
    conn = get_conn()
    conn.execute(
        "INSERT INTO settings (key, value) VALUES (?, ?) "
        "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
        (key, value),
    )
    conn.commit()
    conn.close()


def get_save_folder():
    default = str(Path.home() / "Documents" / "GST Inward HSN Reports")
    folder = get_setting("save_folder", default)
    Path(folder).mkdir(parents=True, exist_ok=True)
    return folder


# ---------- GSTIN management ----------

def list_gstins():
    conn = get_conn()
    rows = conn.execute("SELECT gstin, label FROM gstins ORDER BY gstin").fetchall()
    conn.close()
    return [dict(r) for r in rows]


def add_gstin(gstin, label=""):
    gstin = gstin.strip().upper()
    if not re.match(r"^\d{2}[A-Z0-9]{10}[A-Z0-9]{3}$", gstin):
        raise ValueError(f"'{gstin}' does not look like a valid 15-character GSTIN")
    conn = get_conn()
    conn.execute(
        "INSERT INTO gstins (gstin, label) VALUES (?, ?) "
        "ON CONFLICT(gstin) DO UPDATE SET label=excluded.label",
        (gstin, label),
    )
    conn.commit()
    conn.close()
    return {"gstin": gstin, "label": label}


def remove_gstin(gstin):
    conn = get_conn()
    conn.execute("DELETE FROM gstins WHERE gstin=?", (gstin,))
    conn.commit()
    conn.close()


# ---------- core decode ----------

FILENAME_PATTERN = re.compile(
    r"(?P<gstin>\d{2}[A-Z0-9]{10}[A-Z0-9]{3})_(?P<mm>\d{2})(?P<yyyy>\d{4})_Received",
    re.IGNORECASE,
)
# Looser fallback: just a 6-digit MMYYYY token anywhere in the filename.
LOOSE_MONTH_PATTERN = re.compile(r"(?<!\d)(0[1-9]|1[0-2])(20\d{2})(?!\d)")


def parse_filename(filename):
    """Best-effort extraction of (gstin, period 'YYYY-MM') from a portal-downloaded
    filename like '12ABCDE1234F1Z5_042026_Received_1__1_.json'.
    Returns (gstin_or_None, period_or_None)."""
    name = Path(filename).name
    m = FILENAME_PATTERN.search(name)
    if m:
        mm, yyyy = m.group("mm"), m.group("yyyy")
        return m.group("gstin").upper(), f"{yyyy}-{mm}"
    m2 = LOOSE_MONTH_PATTERN.search(name)
    if m2:
        mm, yyyy = m2.group(1), m2.group(2)
        return None, f"{yyyy}-{mm}"
    return None, None


def _decode_record(rec):
    """Decode one NIC e-invoice 'Received' record into header + item rows."""
    si = rec["SignedInvoice"]
    payload_b64 = si.split(".")[1]
    payload = json.loads(base64.urlsafe_b64decode(_b64pad(payload_b64)))
    inv = json.loads(payload["data"])

    seller = inv.get("SellerDtls", {})
    buyer = inv.get("BuyerDtls", {})
    doc = inv.get("DocDtls", {})
    val = inv.get("ValDtls", {})

    inv_date_raw = doc.get("Dt", "")
    period = None
    for fmt in ("%d/%m/%Y", "%d-%m-%Y"):
        try:
            period = datetime.datetime.strptime(inv_date_raw, fmt).strftime("%Y-%m")
            break
        except ValueError:
            continue
    if period is None:
        # fall back to AckDt
        try:
            period = datetime.datetime.strptime(rec.get("AckDt", ""), "%Y-%m-%d %H:%M:%S").strftime("%Y-%m")
        except Exception:
            period = "unknown"

    header = {
        "irn": inv.get("Irn"),
        "ack_no": rec.get("AckNo"),
        "ack_dt": rec.get("AckDt"),
        "status": rec.get("Status"),
        "doc_type": doc.get("Typ"),
        "invoice_no": doc.get("No"),
        "invoice_date": inv_date_raw,
        "supplier_gstin": seller.get("Gstin"),
        "supplier_name": seller.get("TrdNm", seller.get("LglNm")),
        "buyer_gstin": buyer.get("Gstin"),
        "buyer_name": buyer.get("TrdNm", buyer.get("LglNm")),
        "taxable_value": val.get("AssVal"),
        "igst": val.get("IgstVal"),
        "cgst": val.get("CgstVal"),
        "sgst": val.get("SgstVal"),
        "cess": val.get("CesVal"),
        "total_value": val.get("TotInvVal"),
        "period": period,
    }

    items = []
    for item in inv.get("ItemList", []):
        items.append(
            {
                "irn": header["irn"],
                "sl_no": item.get("SlNo"),
                "hsn": item.get("HsnCd"),
                "description": item.get("PrdDesc"),
                "is_service": item.get("IsServc"),
                "qty": item.get("Qty"),
                "unit": item.get("Unit"),
                "unit_price": item.get("UnitPrice"),
                "taxable_amount": item.get("AssAmt"),
                "gst_rate": item.get("GstRt"),
                "igst_amt": item.get("IgstAmt"),
                "cgst_amt": item.get("CgstAmt"),
                "sgst_amt": item.get("SgstAmt"),
                "cess_amt": item.get("CesAmt"),
                "total_item_value": item.get("TotItemVal"),
            }
        )
    return header, items


def ingest_json_file(json_path, gstin_tag=None, period_override=None):
    """Read a downloaded e-invoice JSON file, decode every record, upsert into DB.
    period_override (e.g. '2026-04'), when given, is used as the period for every
    record in this file, taking priority over each invoice's own date — this
    matches the GST portal's download-batch naming, which is what should drive
    filing-period grouping rather than scattered individual invoice dates.
    Returns a summary dict."""
    with open(json_path, "r", encoding="utf-8") as f:
        raw = json.load(f)

    records = raw if isinstance(raw, list) else raw.get("data", raw.get("Records", []))

    conn = get_conn()
    cur = conn.cursor()
    ok, failed = 0, 0
    periods_seen = set()
    for rec in records:
        try:
            header, items = _decode_record(rec)
        except Exception:
            failed += 1
            continue
        if period_override:
            header["period"] = period_override
        if gstin_tag:
            header["gstin"] = gstin_tag
        else:
            header["gstin"] = header.get("buyer_gstin")

        cur.execute(
            """INSERT INTO invoices (irn, gstin, ack_no, ack_dt, status, doc_type, invoice_no,
               invoice_date, supplier_gstin, supplier_name, buyer_gstin, buyer_name,
               taxable_value, igst, cgst, sgst, cess, total_value, period)
               VALUES (:irn,:gstin,:ack_no,:ack_dt,:status,:doc_type,:invoice_no,
               :invoice_date,:supplier_gstin,:supplier_name,:buyer_gstin,:buyer_name,
               :taxable_value,:igst,:cgst,:sgst,:cess,:total_value,:period)
               ON CONFLICT(irn) DO UPDATE SET
                 status=excluded.status, ack_no=excluded.ack_no, ack_dt=excluded.ack_dt,
                 period=excluded.period, gstin=excluded.gstin""",
            header,
        )
        cur.execute("DELETE FROM items WHERE irn=?", (header["irn"],))
        for it in items:
            cur.execute(
                """INSERT INTO items (irn, sl_no, hsn, description, is_service, qty, unit,
                   unit_price, taxable_amount, gst_rate, igst_amt, cgst_amt, sgst_amt,
                   cess_amt, total_item_value)
                   VALUES (:irn,:sl_no,:hsn,:description,:is_service,:qty,:unit,
                   :unit_price,:taxable_amount,:gst_rate,:igst_amt,:cgst_amt,:sgst_amt,
                   :cess_amt,:total_item_value)""",
                it,
            )
        periods_seen.add(header["period"])
        ok += 1

    conn.commit()
    conn.close()

    return {"ingested": ok, "failed": failed, "periods": sorted(periods_seen)}


# ---------- reporting ----------

def get_available_periods(gstin=None):
    conn = get_conn()
    if gstin:
        rows = conn.execute(
            "SELECT DISTINCT period FROM invoices WHERE gstin=? ORDER BY period", (gstin,)
        ).fetchall()
    else:
        rows = conn.execute("SELECT DISTINCT period FROM invoices ORDER BY period").fetchall()
    conn.close()
    return [r["period"] for r in rows if r["period"] and r["period"] != "unknown"]


def _query_range(period_from, period_to, gstin=None):
    conn = get_conn()
    q = "SELECT * FROM invoices WHERE period BETWEEN ? AND ?"
    params = [period_from, period_to]
    if gstin:
        q += " AND gstin=?"
        params.append(gstin)
    inv_rows = [dict(r) for r in conn.execute(q, params).fetchall()]

    irns = [r["irn"] for r in inv_rows]
    item_rows = []
    if irns:
        placeholders = ",".join("?" * len(irns))
        item_rows = [
            dict(r)
            for r in conn.execute(
                f"SELECT * FROM items WHERE irn IN ({placeholders})", irns
            ).fetchall()
        ]
    conn.close()
    return inv_rows, item_rows


def _query_periods(periods, gstin=None):
    conn = get_conn()
    placeholders = ",".join("?" * len(periods))
    q = f"SELECT * FROM invoices WHERE period IN ({placeholders})"
    params = list(periods)
    if gstin:
        q += " AND gstin=?"
        params.append(gstin)
    inv_rows = [dict(r) for r in conn.execute(q, params).fetchall()]

    irns = [r["irn"] for r in inv_rows]
    item_rows = []
    if irns:
        ph2 = ",".join("?" * len(irns))
        item_rows = [
            dict(r)
            for r in conn.execute(f"SELECT * FROM items WHERE irn IN ({ph2})", irns).fetchall()
        ]
    conn.close()
    return inv_rows, item_rows


def generate_report_for_periods(periods, gstin=None):
    """Consolidate a specific (possibly non-contiguous) set of months, e.g.
    ['2026-01', '2026-03'], into one Excel."""
    periods = sorted(set(periods))
    inv_rows, item_rows = _query_periods(periods, gstin)
    save_folder = get_save_folder()
    tag = gstin if gstin else "AllGSTIN"
    if len(periods) == 1:
        label = periods[0]
    else:
        label = f"{len(periods)}months_{periods[0]}_to_{periods[-1]}"
    fname = f"HSN_Report_{tag}_{label}.xlsx"
    out_path = str(Path(save_folder) / fname)
    _write_excel(inv_rows, item_rows, out_path)

    conn = get_conn()
    conn.execute(
        "INSERT INTO exports (filename, filepath, period_from, period_to, gstin, invoice_count, generated_at) "
        "VALUES (?,?,?,?,?,?,?)",
        (fname, out_path, periods[0], periods[-1], gstin or "", len(inv_rows),
         datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    )
    conn.commit()
    conn.close()
    return {"path": out_path, "invoice_count": len(inv_rows), "item_count": len(item_rows)}


def _write_excel(inv_rows, item_rows, out_path):
    df_inv = pd.DataFrame(inv_rows)
    df_items = pd.DataFrame(item_rows)

    if df_inv.empty:
        df_inv = pd.DataFrame(columns=[
            "irn", "ack_no", "ack_dt", "status", "doc_type", "invoice_no", "invoice_date",
            "supplier_gstin", "supplier_name", "buyer_gstin", "buyer_name",
            "taxable_value", "igst", "cgst", "sgst", "cess", "total_value"])
    else:
        df_inv = df_inv.merge(
            df_items.groupby("irn").size().rename("line_items") if not df_items.empty else pd.Series(dtype=int, name="line_items"),
            left_on="irn", right_index=True, how="left"
        )

    hsn_summary = pd.DataFrame()
    if not df_items.empty:
        active_irns = set(r["irn"] for r in inv_rows if r["status"] == "ACT")
        active_items = df_items[df_items["irn"].isin(active_irns)]
        hsn_summary = active_items.groupby("hsn", as_index=False).agg(
            invoice_count=("irn", "nunique"),
            line_items=("hsn", "count"),
            total_taxable_value=("taxable_amount", "sum"),
            total_igst=("igst_amt", "sum"),
            total_cgst=("cgst_amt", "sum"),
            total_sgst=("sgst_amt", "sum"),
            total_value=("total_item_value", "sum"),
        ).sort_values("total_value", ascending=False)

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        df_inv.to_excel(writer, sheet_name="Invoice Summary", index=False)
        df_items.to_excel(writer, sheet_name="HSN Item Detail", index=False)
        hsn_summary.to_excel(writer, sheet_name="HSN Summary", index=False)

    _format_workbook(out_path)


def _format_workbook(path):
    wb = openpyxl.load_workbook(path)
    header_fill = PatternFill(start_color=HEADER_FILL, end_color=HEADER_FILL, fill_type="solid")
    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    body_font = Font(name="Arial", size=10)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        if ws.max_row < 1:
            continue
        headers = [c.value for c in ws[1]]
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 26
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
        for i, col_name in enumerate(headers, start=1):
            letter = get_column_letter(i)
            sample_rows = range(2, min(ws.max_row, 200) + 1)
            maxlen = max(
                [len(str(col_name))]
                + [len(str(ws.cell(row=r, column=i).value)) for r in sample_rows if ws.cell(row=r, column=i).value is not None]
                or [10]
            )
            ws.column_dimensions[letter].width = min(max(maxlen + 2, 10), 40)
        if ws.max_row >= 1 and ws.max_column >= 1:
            ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def generate_report(period_from, period_to, gstin=None):
    inv_rows, item_rows = _query_range(period_from, period_to, gstin)
    save_folder = get_save_folder()
    tag = gstin if gstin else "AllGSTIN"
    fname = f"HSN_Report_{tag}_{period_from}_to_{period_to}.xlsx"
    out_path = str(Path(save_folder) / fname)
    _write_excel(inv_rows, item_rows, out_path)

    conn = get_conn()
    conn.execute(
        "INSERT INTO exports (filename, filepath, period_from, period_to, gstin, invoice_count, generated_at) "
        "VALUES (?,?,?,?,?,?,?)",
        (fname, out_path, period_from, period_to, gstin or "", len(inv_rows),
         datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    )
    conn.commit()
    conn.close()
    return {"path": out_path, "invoice_count": len(inv_rows), "item_count": len(item_rows)}


def clear_all_data():
    """Wipe all ingested invoice/item/export history. GSTIN registrations and
    the save-folder setting are left intact — only uploaded invoice data and
    export history are cleared. This cannot be undone from within the app;
    the only way to get the data back is to re-upload the original JSON
    file(s)."""
    conn = get_conn()
    conn.execute("DELETE FROM invoices")
    conn.execute("DELETE FROM items")
    conn.execute("DELETE FROM exports")
    conn.commit()
    conn.close()


def list_exports():
    conn = get_conn()
    rows = conn.execute("SELECT * FROM exports ORDER BY id DESC LIMIT 50").fetchall()
    conn.close()
    return [dict(r) for r in rows]
